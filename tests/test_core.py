"""Unit tests for parsing, OCR reconciliation and master-sheet upsert (no OCR/browser needed).

Run:  python -m pytest -q
"""

from datetime import datetime

from openpyxl import load_workbook

from txdot_bid_tabs.extractor import RawRow, normalize_item, parse_money, parse_qty, reconcile
from txdot_bid_tabs.master_sheet import upsert_project
from txdot_bid_tabs.models import Bid, BidItem, Project


def test_parse_money():
    assert parse_money("$150,000.00") == 150000.0
    assert parse_money("$1.50") == 1.5
    assert parse_money("$15000000") == 150000.0   # OCR dropped the decimal point
    assert parse_money("") is None


def test_parse_qty():
    assert parse_qty("18,750") == 18750
    assert parse_qty("") is None


def test_normalize_item():
    assert normalize_item("F-164,2") == "F-164.2"
    assert normalize_item("F-164.3C") == "F-164.3c"
    assert normalize_item("tx-100") == "TX-100"
    assert normalize_item("C105") == "C-105"


def test_reconcile_fixes_misread_total():
    # Real misread seen on the Lockhart scan: 3 x $3,200 printed as $9,660.
    rows = [RawRow("C-105", 1, "LS", "MOBILIZATION", 150000, 150000),
            RawRow("F-164.3b", 3, "EA", "SWING GATE", 3200, 9660)]
    up, tot, verified, notes = reconcile(rows, base_bid=159600)
    assert verified
    assert tot["F-164.3b"] == 9600 and up["F-164.3b"] == 3200
    assert notes  # the correction is logged


def test_reconcile_fixes_misread_unit_price():
    rows = [RawRow("F-164.2", 17730, "LF", "FENCING", 84.0, 602820)]  # 34 misread as 84
    up, tot, verified, _ = reconcile(rows, base_bid=602820)
    assert verified and up["F-164.2"] == 34.0


def test_reconcile_flags_unverifiable():
    rows = [RawRow("C-105", 1, "LS", "MOBILIZATION", 100, 100)]
    _, _, verified, notes = reconcile(rows, base_bid=999)
    assert not verified and any("NOT verified" in n for n in notes)


def _project(pid, price):
    items = [BidItem("C-105", "MOBILIZATION", "LS", 1)]
    bids = [Bid("Low Co", price, {"C-105": price}, {"C-105": price}, True),
            Bid("High Co", price * 2, {"C-105": price * 2}, {"C-105": price * 2}, True)]
    return Project(pid, "Test Municipal Airport", "1234TEST", "Paving", "", datetime(2026, 10, 1),
                   items=items, bids=bids)


def test_master_upsert_has_no_duplicates(tmp_path):
    path = tmp_path / "master.xlsx"
    upsert_project(_project("NTB1", 100.0), path)
    upsert_project(_project("NTB1", 120.0), path)   # re-run same project with new data
    upsert_project(_project("NTB2", 200.0), path)
    wb = load_workbook(path)
    data = list(wb["Bid Data"].iter_rows(min_row=2, values_only=True))
    assert len(data) == 4                                  # 2 projects x 2 bidders x 1 item
    assert {r[0] for r in data} == {"NTB1", "NTB2"}
    ntb1_low = [r for r in data if r[0] == "NTB1" and r[6] == "Yes"][0]
    assert ntb1_low[11] == 120.0                            # replaced, not appended
    assert wb["Projects"].max_row == 3


def test_master_preserves_region(tmp_path):
    path = tmp_path / "master.xlsx"
    upsert_project(_project("NTB1", 100.0), path)
    wb = load_workbook(path)
    ws = wb["Projects"]
    region_col = [c.value for c in ws[1]].index("Region") + 1
    ws.cell(row=2, column=region_col, value="Central")     # user fills in a region
    wb.save(path)
    upsert_project(_project("NTB1", 110.0), path)
    ws = load_workbook(path)["Projects"]
    assert ws.cell(row=2, column=region_col).value == "Central"
