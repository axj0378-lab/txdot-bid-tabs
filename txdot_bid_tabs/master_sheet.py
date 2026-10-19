"""Stage 6: the statewide master unit-price workbook ("living database").

Three sheets:
  * Bid Data            the database: one row per bidder x item, for EVERY bidder
                        (low bid flagged). Slide 12 asks for low bids only for now,
                        but keeping all bids costs nothing and enables the
                        all-bids analysis it mentions for later.
  * Projects            one row per project: dates, low bidder, Region (blank,
                        to be filled by the regionalization step), and a
                        live "Price Age" bucket (last quarter / 6 months / year / older).
  * Master Unit Prices  the view from slides 12-13: rows are bid items, each
                        project adds a Unit Price / Qty column block with its
                        project info above it. Cells are SUMIFS formulas over
                        Bid Data (low bid only), and the right-hand columns
                        average the low-bid unit price by recency, the first
                        estimating aid from slide 14.

Re-running for a project already in the workbook replaces that project's rows
(upsert), so the workbook never holds duplicates. Values typed into the
Projects "Region" column are preserved across runs.
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import MASTER_SHEET_PATH
from .models import Project

FONT = "Arial"
NAVY = "1F3864"
HEAD_FILL = "D9E1F2"
INFO_FILL = "F2F2F2"
MONEY = '$#,##0.00;($#,##0.00);"-"'
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

DATA_COLS = ["Project ID", "Airport", "CSJ", "Bid Date", "Bidder", "Bid Rank", "Low Bid",
             "Item", "Description", "Unit", "Qty", "Unit Price", "Total", "Verified"]
PROJECT_COLS = ["Project ID", "Airport", "CSJ", "Scope", "Description", "Bid Date",
                "Location", "Region", "Bidders", "Low Bidder", "Low Bid", "Price Age",
                "Source PDF", "Last Updated"]

# Price Age buckets (slide 12), in days
AGE_BUCKETS = [(91, "Last quarter"), (183, "Last 6 months"), (365, "Last year")]


def _age_formula(date_ref: str) -> str:
    f = '"Older than 1 year"'
    for days, label in reversed(AGE_BUCKETS):
        f = f'IF(TODAY()-{date_ref}<={days},"{label}",{f})'
    return f'=IF({date_ref}="","",{f})'


def _item_sort_key(code: str):
    """Natural sort: C-105 < F-164.1 < F-164.3a < L-125 < TX-100."""
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", code)]


# --------------------------------------------------------------------------
# Load / upsert
# --------------------------------------------------------------------------

def _read_sheet(path: Path, name: str) -> list[dict]:
    if not path.exists():
        return []
    wb = load_workbook(path, data_only=True)
    if name not in wb.sheetnames:
        return []
    rows = list(wb[name].iter_rows(values_only=True))
    if not rows:
        return []
    header = list(rows[0])
    return [dict(zip(header, r)) for r in rows[1:] if any(v is not None for v in r)]


def _project_rows(project: Project, source_url: str) -> tuple[list[dict], dict]:
    low = project.low_bid
    ranked = sorted(project.bids, key=lambda b: b.base_bid)
    data = []
    for rank, b in enumerate(ranked, start=1):
        for it in project.items:
            data.append({
                "Project ID": project.project_id, "Airport": project.airport, "CSJ": project.csj,
                "Bid Date": project.bid_date, "Bidder": b.bidder, "Bid Rank": rank,
                "Low Bid": "Yes" if b is low else "No", "Item": it.item,
                "Description": it.description, "Unit": it.unit, "Qty": it.qty,
                "Unit Price": b.unit_prices.get(it.item), "Total": b.totals.get(it.item),
                "Verified": "Yes" if b.verified else "No",
            })
    proj = {
        "Project ID": project.project_id, "Airport": project.airport, "CSJ": project.csj,
        "Scope": project.scope, "Description": project.description,
        "Bid Date": project.bid_date, "Location": project.location, "Region": None,
        "Bidders": len(project.bids), "Low Bidder": low.bidder, "Low Bid": low.base_bid,
        "Source PDF": source_url, "Last Updated": datetime.now().replace(microsecond=0),
    }
    return data, proj


def upsert_project(project: Project, path: Path = MASTER_SHEET_PATH, source_url: str = "") -> Path:
    """Add or replace one project in the master workbook, then rebuild its views."""
    data = [r for r in _read_sheet(path, "Bid Data") if r.get("Project ID") != project.project_id]
    old_projects = _read_sheet(path, "Projects")
    regions = {p["Project ID"]: p.get("Region") for p in old_projects}
    projects = [p for p in old_projects if p.get("Project ID") != project.project_id]

    new_data, new_proj = _project_rows(project, source_url)
    new_proj["Region"] = regions.get(project.project_id)
    data += new_data
    projects.append(new_proj)
    projects.sort(key=lambda p: p["Bid Date"])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    _write_master_view(wb, data, projects)
    _write_table(wb, "Projects", PROJECT_COLS, projects)
    _write_table(wb, "Bid Data", DATA_COLS, data)
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path


# --------------------------------------------------------------------------
# Sheets
# --------------------------------------------------------------------------

def _style(c, *, bold=False, fill=None, fmt=None, align=None, border=False, size=10, color="000000",
           wrap=False):
    c.font = Font(name=FONT, bold=bold, size=size, color=color)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if border:
        c.border = BOX
    return c


def _write_table(wb: Workbook, name: str, cols: list[str], rows: list[dict]) -> None:
    ws = wb.create_sheet(name)
    for j, h in enumerate(cols, start=1):
        _style(ws.cell(row=1, column=j, value=h), bold=True, fill=HEAD_FILL, border=True)
    for i, row in enumerate(rows, start=2):
        for j, h in enumerate(cols, start=1):
            c = ws.cell(row=i, column=j, value=row.get(h))
            _style(c)
            if h in ("Bid Date", "Last Updated"):
                c.number_format = "yyyy-mm-dd hh:mm" if h == "Last Updated" else "yyyy-mm-dd"
            elif h in ("Unit Price", "Total", "Low Bid"):
                c.number_format = MONEY
            elif h == "Qty":
                c.number_format = "#,##0.##"
        if name == "Projects":
            date_col = get_column_letter(cols.index("Bid Date") + 1)
            _style(ws.cell(row=i, column=cols.index("Price Age") + 1,
                           value=_age_formula(f"{date_col}{i}")))
            region = ws.cell(row=i, column=cols.index("Region") + 1)
            region.fill = PatternFill("solid", fgColor="FFF2CC")  # user-editable input
    widths = {"Project ID": 20, "Airport": 30, "Bidder": 32, "Description": 44, "Low Bidder": 32,
              "Source PDF": 60, "Price Age": 18, "Last Updated": 17, "Scope": 14, "Location": 14}
    for j, h in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(j)].width = widths.get(h, 12)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(len(rows) + 1, 2)}"


def _write_master_view(wb: Workbook, data: list[dict], projects: list[dict]) -> None:
    ws = wb.create_sheet("Master Unit Prices")
    ws.sheet_view.showGridLines = False
    D = "'Bid Data'!"
    col = {h: f"{D}${get_column_letter(i + 1)}:${get_column_letter(i + 1)}"
           for i, h in enumerate(DATA_COLS)}
    P = "Projects!"
    pcol = {h: f"{P}${get_column_letter(i + 1)}:${get_column_letter(i + 1)}"
            for i, h in enumerate(PROJECT_COLS)}

    _style(ws.cell(row=1, column=1, value="Statewide Unit Bid Prices (Low Bid)"),
           bold=True, size=14, color=NAVY)
    _style(ws.cell(row=2, column=1, value="One row per bid item; each project adds a Unit Price / Qty "
                                          "block. Values are live formulas over the Bid Data sheet."),
           color="595959")

    # Items: union across projects (low-bid rows), keyed by item code + unit;
    # description from the most recent project that used the item.
    items: dict[tuple, dict] = {}
    for r in sorted((r for r in data if r["Low Bid"] == "Yes"), key=lambda r: r["Bid Date"]):
        items[(r["Item"], r["Unit"])] = r
    item_keys = sorted(items, key=lambda k: _item_sort_key(k[0]))

    info_labels = ["Project", "Project ID", "CSJ", "Bid Date", "Region", "Price Age", "Low Bidder"]
    info_top = 4
    hdr = info_top + len(info_labels)
    first = hdr + 1

    for k, label in enumerate(info_labels):
        r = info_top + k
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        _style(ws.cell(row=r, column=1, value=label), bold=True, align="right")
    for j, h in enumerate(["Item", "Description", "Unit"], start=1):
        _style(ws.cell(row=hdr, column=j, value=h), bold=True, fill=HEAD_FILL, align="center",
               border=True)

    for p_idx, p in enumerate(projects):
        c_up = 4 + 2 * p_idx
        c_qty = c_up + 1
        UP, QT = get_column_letter(c_up), get_column_letter(c_qty)
        pid_ref = f"{UP}${info_top + 1}"
        info_values = [
            f"{p['Airport'].replace(' Airport', '')} ({p['Bid Date']:%B %Y})",
            p["Project ID"],
            f"=INDEX({pcol['CSJ']},MATCH({pid_ref},{pcol['Project ID']},0))",
            f"=INDEX({pcol['Bid Date']},MATCH({pid_ref},{pcol['Project ID']},0))",
            f'=IF(INDEX({pcol["Region"]},MATCH({pid_ref},{pcol["Project ID"]},0))="","(not set)",'
            f'INDEX({pcol["Region"]},MATCH({pid_ref},{pcol["Project ID"]},0)))',
            _age_formula(f"{UP}${info_top + 3}"),
            f"=INDEX({pcol['Low Bidder']},MATCH({pid_ref},{pcol['Project ID']},0))",
        ]
        for k, v in enumerate(info_values):
            r = info_top + k
            ws.merge_cells(start_row=r, start_column=c_up, end_row=r, end_column=c_qty)
            _style(ws.cell(row=r, column=c_up, value=v), bold=(k == 0), fill=INFO_FILL,
                   align="center", border=True, fmt="mmm d, yyyy" if k == 3 else None, wrap=True)
            ws.cell(row=r, column=c_qty).border = BOX
        for k in (0, len(info_labels) - 1):  # project name + low bidder can be long
            ws.row_dimensions[info_top + k].height = 30
        _style(ws.cell(row=hdr, column=c_up, value="Unit Price"), bold=True, fill=HEAD_FILL,
               align="center", border=True)
        _style(ws.cell(row=hdr, column=c_qty, value="Qty"), bold=True, fill=HEAD_FILL,
               align="center", border=True)
        for i, key in enumerate(item_keys):
            r = first + i
            crit = (f"{col['Project ID']},{pid_ref},{col['Item']},$A{r},"
                    f"{col['Unit']},$C{r},{col['Low Bid']},\"Yes\"")
            # blank (not $0) when this project didn't bid the item
            ws.cell(row=r, column=c_up,
                    value=f'=IF(COUNTIFS({crit})=0,"",SUMIFS({col["Unit Price"]},{crit}))')
            ws.cell(row=r, column=c_qty,
                    value=f'=IF(COUNTIFS({crit})=0,"",SUMIFS({col["Qty"]},{crit}))')
            _style(ws.cell(row=r, column=c_up), fmt=MONEY, border=True)
            _style(ws.cell(row=r, column=c_qty), fmt="#,##0.##", align="center", border=True)
        ws.column_dimensions[UP].width = 18
        ws.column_dimensions[QT].width = 12

    # Estimating aids: average low-bid unit price by recency (slide 14)
    est_first = 4 + 2 * len(projects) + 1
    windows = [("Last Quarter", 91), ("Last 6 Months", 183), ("Last Year", 365), ("All Projects", None)]
    ws.merge_cells(start_row=hdr - 1, start_column=est_first, end_row=hdr - 1,
                   end_column=est_first + len(windows))
    _style(ws.cell(row=hdr - 1, column=est_first, value="Average Low-Bid Unit Price (not inflation-adjusted)"),
           bold=True, fill=INFO_FILL, align="center", border=True)
    for k, (label, days) in enumerate(windows):
        c = est_first + k
        _style(ws.cell(row=hdr, column=c, value=label), bold=True, fill=HEAD_FILL, align="center",
               border=True, wrap=True)
        ws.column_dimensions[get_column_letter(c)].width = 14
        for i in range(len(item_keys)):
            r = first + i
            crit = f"{col['Item']},$A{r},{col['Unit']},$C{r},{col['Low Bid']},\"Yes\""
            if days:
                crit += f",{col['Bid Date']},\">=\"&(TODAY()-{days})"
            _style(ws.cell(row=r, column=c,
                           value=f'=IFERROR(AVERAGEIFS({col["Unit Price"]},{crit}),"")'),
                   fmt=MONEY, border=True)
    c = est_first + len(windows)
    _style(ws.cell(row=hdr, column=c, value="# Projects"), bold=True, fill=HEAD_FILL,
           align="center", border=True)
    ws.column_dimensions[get_column_letter(c)].width = 11
    for i in range(len(item_keys)):
        r = first + i
        _style(ws.cell(row=r, column=c,
                       value=f"=COUNTIFS({col['Item']},$A{r},{col['Unit']},$C{r},{col['Low Bid']},\"Yes\")"),
               align="center", border=True)

    for i, key in enumerate(item_keys):
        r = first + i
        it = items[key]
        _style(ws.cell(row=r, column=1, value=it["Item"]), align="center", border=True)
        _style(ws.cell(row=r, column=2, value=it["Description"]), border=True, wrap=True)
        _style(ws.cell(row=r, column=3, value=it["Unit"]), align="center", border=True)
    ws.column_dimensions["A"].width = 11
    ws.column_dimensions["B"].width = 44
    ws.column_dimensions["C"].width = 7
    ws.freeze_panes = ws.cell(row=first, column=4)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
