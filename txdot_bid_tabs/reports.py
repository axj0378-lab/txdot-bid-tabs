"""Stages 4-5: per-project bid summary (bid tabulation) as Excel and PDF.

Layout (slides 10-11): item columns (Item, Qty, Unit, Description) followed by a
Unit Price / Total column pair per bidder with the company name above each pair.
Bidders are ordered low -> high so the low bidder is the first block, and it is
highlighted. Footer rows show each bidder's total, the base bid typed on TxDOT's
summary sheet, the difference (a built-in check that should read $0), rank, and
% above the low bid.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .config import OUTPUT_DIR
from .models import Project

MONEY = '$#,##0.00;($#,##0.00);"-"'
QTY = "#,##0.##"

FONT = "Arial"
NAVY = "1F3864"
LOW_FILL = "E2EFDA"   # light green: low bidder
HEAD_FILL = "D9E1F2"
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _ranked_bids(project: Project):
    return sorted(project.bids, key=lambda b: b.base_bid)


def _file_stem(project: Project) -> str:
    return f"{project.project_id}_bid_summary"


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------

def write_summary_excel(project: Project, out_dir: Path = OUTPUT_DIR, source_url: str = "") -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    bids = _ranked_bids(project)
    low = project.low_bid
    n_items = len(project.items)

    wb = Workbook()
    ws = wb.active
    ws.title = "Bid Summary"
    ws.sheet_view.showGridLines = False

    def cell(r, c, value=None, *, bold=False, fill=None, fmt=None, align=None, size=10,
             color="000000", border=False, wrap=False):
        x = ws.cell(row=r, column=c, value=value)
        x.font = Font(name=FONT, bold=bold, size=size, color=color)
        if fill:
            x.fill = PatternFill("solid", fgColor=fill)
        if fmt:
            x.number_format = fmt
        x.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
        if border:
            x.border = BOX
        return x

    # Title + project information block
    cell(1, 1, f"Bid Tabulation: {project.airport}", bold=True, size=14, color=NAVY)
    info = [
        ("Project ID", project.project_id),
        ("TxDOT CSJ #", project.csj),
        ("Project Scope", project.scope),
        ("Description", project.description),
        ("Bid Opening", f"{project.bid_date:%B %d, %Y %I:%M %p}".replace(" 0", " ")),
        ("Bid Opening Location", project.location),
        ("Number of Bidders", len(bids)),
        ("Low Bidder", f"{low.bidder}  ({low.base_bid:,.2f})".replace("(", "($")),
    ]
    for i, (k, v) in enumerate(info, start=3):
        cell(i, 1, k, bold=True)
        cell(i, 2, v, bold=(k == "Low Bidder"), align="left")

    # Table header (two rows): bidder names, then column names
    top = 3 + len(info) + 1
    hdr = top + 1
    first_data = hdr + 1
    last_data = first_data + n_items - 1
    fixed = ["Item", "Qty", "Unit", "Description"]
    for c, name in enumerate(fixed, start=1):
        ws.merge_cells(start_row=top, start_column=c, end_row=hdr, end_column=c)
        cell(top, c, name, bold=True, fill=HEAD_FILL, align="center", border=True)
        ws.cell(row=hdr, column=c).border = BOX

    bidder_cols = []  # (unit price col, total col)
    for i, b in enumerate(bids):
        up_col = len(fixed) + 1 + 2 * i
        tot_col = up_col + 1
        bidder_cols.append((up_col, tot_col))
        is_low = b is low
        fill = LOW_FILL if is_low else HEAD_FILL
        ws.merge_cells(start_row=top, start_column=up_col, end_row=top, end_column=tot_col)
        label = f"{b.bidder}\nLOW BIDDER" if is_low else b.bidder
        cell(top, up_col, label, bold=True, fill=fill, align="center", border=True, wrap=True)
        ws.cell(row=top, column=tot_col).border = BOX
        cell(hdr, up_col, "Unit Price", bold=True, fill=fill, align="center", border=True)
        cell(hdr, tot_col, "Total", bold=True, fill=fill, align="center", border=True)
    ws.row_dimensions[top].height = 32

    # Line items. Totals are formulas (Qty x Unit Price) so the sheet stays live.
    for r, item in enumerate(project.items, start=first_data):
        cell(r, 1, item.item, align="center", border=True)
        cell(r, 2, item.qty, fmt=QTY, align="center", border=True)
        cell(r, 3, item.unit, align="center", border=True)
        cell(r, 4, item.description, border=True, wrap=True)
        for b, (up_col, tot_col) in zip(bids, bidder_cols):
            fill = LOW_FILL if b is low else None
            cell(r, up_col, b.unit_prices.get(item.item), fmt=MONEY, fill=fill, border=True)
            cell(r, tot_col, f"=B{r}*{get_column_letter(up_col)}{r}", fmt=MONEY, fill=fill, border=True)

    # Footer rows
    r_total, r_base, r_diff, r_rank, r_pct = (last_data + 1 + k for k in range(5))
    labels = {
        r_total: "Total Bid (sum of line items)",
        r_base: "Base Bid per TxDOT Summary Sheet",
        r_diff: "Difference (should be $0)",
        r_rank: "Rank",
        r_pct: "Above Low Bid",
    }
    for r, text in labels.items():
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        cell(r, 1, text, bold=True, align="right")
    total_refs = [f"{get_column_letter(t)}{r_total}" for _, t in bidder_cols]
    for b, (up_col, tot_col) in zip(bids, bidder_cols):
        L = get_column_letter(tot_col)
        fill = LOW_FILL if b is low else None
        low_ref = f"${get_column_letter(bidder_cols[0][1])}${r_total}"  # first block = low bidder
        cell(r_total, tot_col, f"=SUM({L}{first_data}:{L}{last_data})", bold=True, fmt=MONEY,
             fill=fill, border=True)
        cell(r_base, tot_col, b.base_bid, fmt=MONEY, fill=fill, border=True)
        cell(r_diff, tot_col, f"={L}{r_total}-{L}{r_base}", fmt=MONEY, fill=fill, border=True)
        # rank = 1 + number of bidders with a lower total (bid totals sit in
        # non-adjacent columns, so RANK() over a range doesn't apply)
        me = f"{L}{r_total}"
        others = "".join(f"+({me}>{ref})" for ref in total_refs if ref != me)
        cell(r_rank, tot_col, f"=1{others}", fmt="0", fill=fill, align="center", border=True)
        cell(r_pct, tot_col, f"={L}{r_total}/{low_ref}-1", fmt='0.0%;-0.0%;"-"', fill=fill,
             align="center", border=True)

    note_row = r_pct + 2
    cell(note_row, 1, "Notes", bold=True)
    cell(note_row + 1, 1, "Bidders are ordered low to high by base bid; the low bidder is shaded green.")
    cell(note_row + 2, 1, "Unit prices were read from scanned bid tabs (OCR) and verified: every bidder's "
                          "line totals reconcile to the base bid on TxDOT's summary sheet."
         if all(b.verified for b in bids) else
         "WARNING: one or more bidders could not be verified against the summary sheet; see Extraction Log.")
    cell(note_row + 3, 1, f"Source: {source_url or 'TxDOT Aviation Division bid results'}. "
                          f"Generated {datetime.now():%Y-%m-%d %H:%M}.", color="595959")

    widths = {1: 22, 2: 10, 3: 7, 4: 44}
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    for up_col, tot_col in bidder_cols:
        ws.column_dimensions[get_column_letter(up_col)].width = 15
        ws.column_dimensions[get_column_letter(tot_col)].width = 16
    ws.freeze_panes = ws.cell(row=first_data, column=5)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    _write_log_sheet(wb, project)
    wb.calculation.fullCalcOnLoad = True
    path = out_dir / f"{_file_stem(project)}.xlsx"
    wb.save(path)
    return path


def _write_log_sheet(wb: Workbook, project: Project) -> None:
    """Audit trail: what was verified and what the extractor corrected."""
    ws = wb.create_sheet("Extraction Log")
    ws.append(["Bidder", "Base Bid (summary sheet)", "Sum of OCR line totals", "Verified"])
    for b in project.bids:
        ws.append([b.bidder, b.base_bid, sum(b.totals.values()), "Yes" if b.verified else "NO"])
    ws.append([])
    ws.append(["Extraction notes"])
    for n in project.notes or ["No corrections were needed."]:
        ws.append([n])
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name=FONT, bold=c.row == 1 or c.value == "Extraction notes")
            if c.column in (2, 3) and isinstance(c.value, (int, float)):
                c.number_format = MONEY
    ws.column_dimensions["A"].width = 40
    for col in "BCD":
        ws.column_dimensions[col].width = 24


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def write_summary_pdf(project: Project, out_dir: Path = OUTPUT_DIR, source_url: str = "") -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    bids = _ranked_bids(project)
    low = project.low_bid
    path = out_dir / f"{_file_stem(project)}.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=landscape(letter),
                            leftMargin=0.4 * inch, rightMargin=0.4 * inch,
                            topMargin=0.5 * inch, bottomMargin=0.5 * inch,
                            title=f"Bid Tabulation - {project.airport}")

    navy = colors.HexColor("#" + NAVY)
    green = colors.HexColor("#" + LOW_FILL)
    head = colors.HexColor("#" + HEAD_FILL)
    s_title = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=15, textColor=navy, spaceAfter=4)
    s_small = ParagraphStyle("s", fontName="Helvetica", fontSize=7.5, leading=9)
    s_cell = ParagraphStyle("c", fontName="Helvetica", fontSize=7, leading=8.5)
    s_head = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=7.5, leading=9, alignment=1)

    story = [Paragraph(f"Bid Tabulation: {project.airport}", s_title)]
    info = [
        ["Project ID", project.project_id, "Bid Opening",
         f"{project.bid_date:%B %d, %Y %I:%M %p}".replace(" 0", " ")],
        ["TxDOT CSJ #", project.csj, "Location", project.location],
        ["Scope", project.scope, "Bidders", str(len(bids))],
        ["Description", project.description, "Low Bidder", f"{low.bidder} (${low.base_bid:,.2f})"],
    ]
    t = Table(info, colWidths=[1.0 * inch, 3.2 * inch, 1.0 * inch, 4.0 * inch], hAlign="LEFT")
    t.setStyle(TableStyle([
        ("FONT", (0, 0), (-1, -1), "Helvetica", 8.5),
        ("FONT", (0, 0), (0, -1), "Helvetica-Bold", 8.5),
        ("FONT", (2, 0), (2, -1), "Helvetica-Bold", 8.5),
        ("FONT", (3, 3), (3, 3), "Helvetica-Bold", 8.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5), ("TOPPADDING", (0, 0), (-1, -1), 1.5),
    ]))
    story += [t, Spacer(1, 8)]

    # Comparison table
    nb = len(bids)
    row0 = ["Item", "Qty", "Unit", "Description"]
    row1 = ["", "", "", ""]
    for b in bids:
        name = b.bidder + ("<br/><font color='#375623'>LOW BIDDER</font>" if b is low else "")
        row0 += [Paragraph(name, s_head), ""]
        row1 += ["Unit Price", "Total"]
    data = [row0, row1]
    for it in project.items:
        r = [it.item, f"{it.qty:,.0f}" if it.qty == int(it.qty) else f"{it.qty:,.2f}", it.unit,
             Paragraph(it.description, s_cell)]
        for b in bids:
            up = b.unit_prices.get(it.item)
            tot = b.totals.get(it.item)
            r += [f"${up:,.2f}" if up is not None else "-", f"${tot:,.2f}" if tot is not None else "-"]
        data.append(r)

    totals = [sum(b.totals.values()) for b in bids]
    foot = [
        ("Total Bid", [f"${t:,.2f}" for t in totals]),
        ("Base Bid per TxDOT Summary Sheet", [f"${b.base_bid:,.2f}" for b in bids]),
        ("Rank", [str(i + 1) for i in range(nb)]),
        ("Above Low Bid", ["-"] + [f"{t / totals[0] - 1:.1%}" for t in totals[1:]]),
    ]
    first_foot = len(data)
    for label, vals in foot:
        r = [label, "", "", ""]
        for v in vals:
            r += ["", v]
        data.append(r)

    avail = 10.2 * inch
    fixed_w = [0.7 * inch, 0.55 * inch, 0.4 * inch]
    per_bid = min(1.9 * inch, (avail - sum(fixed_w) - 2.0 * inch) / nb)
    desc_w = avail - sum(fixed_w) - per_bid * nb
    widths = fixed_w + [desc_w] + [per_bid * 0.47, per_bid * 0.53] * nb

    style = [
        ("FONT", (0, 0), (-1, -1), "Helvetica", 7.5),
        ("FONT", (0, 0), (-1, 1), "Helvetica-Bold", 7.5),
        ("BACKGROUND", (0, 0), (-1, 1), head),
        ("GRID", (0, 0), (-1, first_foot - 1), 0.4, colors.HexColor("#A6A6A6")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (2, -1), "CENTER"),
        ("ALIGN", (4, 2), (-1, -1), "RIGHT"),
        ("ALIGN", (4, 1), (-1, 1), "CENTER"),
        ("SPAN", (0, 0), (0, 1)), ("SPAN", (1, 0), (1, 1)),
        ("SPAN", (2, 0), (2, 1)), ("SPAN", (3, 0), (3, 1)),
        ("LINEABOVE", (0, first_foot), (-1, first_foot), 1.0, navy),
        ("FONT", (0, first_foot), (-1, first_foot), "Helvetica-Bold", 7.5),
    ]
    for i in range(nb):
        c0 = 4 + 2 * i
        style.append(("SPAN", (c0, 0), (c0 + 1, 0)))
        if bids[i] is low:
            style.append(("BACKGROUND", (c0, 0), (c0 + 1, -1), green))
    for k in range(len(foot)):
        r = first_foot + k
        style += [("SPAN", (0, r), (3, r)), ("ALIGN", (0, r), (3, r), "RIGHT"),
                  ("FONT", (0, r), (3, r), "Helvetica-Bold", 7.5)]
    tbl = Table(data, colWidths=widths, repeatRows=2)
    tbl.setStyle(TableStyle(style))
    story += [tbl, Spacer(1, 10)]

    verified = all(b.verified for b in bids)
    story.append(Paragraph(
        "Bidders are ordered low to high by base bid; the low bidder is shaded green. "
        + ("Unit prices were read from TxDOT's scanned bid tabs and verified: each bidder's line "
           "totals reconcile exactly to the base bid on the summary sheet."
           if verified else
           "<b>WARNING:</b> one or more bidders could not be verified against the summary sheet."),
        s_small))
    story.append(Paragraph(
        f"Source: {source_url or 'TxDOT Aviation Division bid results'}. "
        f"Generated {datetime.now():%B %d, %Y}.", s_small))
    doc.build(story)
    return path
